"""Prepara i fogli di annotazione manuale (gold standard) per il campione di campione_annotazione.py.
Pre-compila con l'output AI esistente (gpt-4.1); gli annotatori correggono.
Non include nulla dal questionario GINO (annotazione in cieco).
Uso: python codice/campione_annotazione.py; python codice/prepara_annotazione.py"""
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from config import MICRODATI, RIFERIMENTI, REPO

OUT = REPO / "annotazione"
GIUDIZI = "ok,non_ente,duplicato,fuori_rete,generico"

ana = pd.read_excel(MICRODATI / "rif_anagrafica_v1.xlsx", dtype=str)
camp = pd.read_csv(OUT / "campione_v1.csv", sep=";", encoding="utf-8-sig", dtype=str)
ASS = dict(zip(camp.codice, camp.annotatori))
sel = ana[ana.codice.isin(camp.codice)].sort_values(["regione", "codice"])
cat = pd.read_csv(RIFERIMENTI / "categorie_gino_30.csv", sep=";", encoding="utf-8-sig")
mp = pd.read_csv(RIFERIMENTI / "rif_mapping_tipi_ai_gino_v1.csv", sep=";", encoding="utf-8-sig")
MAP = dict(zip(mp.tipo_AI, mp.codice))

righe, prot = [], []
for _, r in sel.iterrows():
    c = r.codice
    d = json.load(open(MICRODATI / "ai" / f"ai_{c}.json"))
    testo = (MICRODATI / "protocolli" / f"protocollo_{c}.txt").read_text(encoding="utf-8", errors="ignore")
    sogg = d.get("soggetti", [])
    prot.append(dict(codice=c, regione=r.regione, annotatori=ASS[c], titolo_documento=r.titolo_documento,
                     file_originale=r.allegato_gino_3_9, caratteri=len(testo), enti_AI=len(sogg)))
    for s in sogg:
        t = s.get("tipo", "")
        righe.append(dict(codice=c, regione=r.regione, annotatori=ASS[c], nome_AI=s.get("nome", ""), tipo_AI=t,
                          cat_AI=MAP.get(t, ""), giudizio="", nome_corretto="", cat_principale="",
                          cat_secondarie="", note=""))
PROT = pd.DataFrame(prot)
ENTI = pd.DataFrame(righe)
ENTI.insert(0, "id", range(1, len(ENTI) + 1))  # id globale, uguale in tutti i file

ISTR = [
    "Per ogni protocollo (foglio 'protocolli'), leggi il testo originale e controlla le righe di 'enti_AI' con quel codice.",
    "Colonna 'giudizio': ok = ente corretto; non_ente = non è un ente (norma, ministero solo citato, ecc.); duplicato = stesso ente di un'altra riga; fuori_rete = esiste ma non partecipa alla rete; generico = citazione generica (es. 'i Comuni dell'ambito').",
    "Colonna 'nome_corretto': solo se il nome AI è sbagliato o incompleto.",
    "'cat_principale': un codice da 1 a 30 (foglio 'categorie'). 'cat_secondarie': altri codici separati da ';'.",
    "Enti presenti nel protocollo e NON nella lista AI: aggiungili nel foglio 'enti_mancanti'.",
    "Regole di classificazione: docs/definizioni_soggetti_v1.md. Nel dubbio usa 30 (Altro) e scrivi una nota.",
    "La colonna 'cat_AI' è la categoria proposta dall'AI (mappata sui 30 codici): non è la verità, va controllata.",
    "Colonna 'annotatori': A1+A2 = protocollo annotato da entrambi; A1 o A2 = annotato solo da uno.",
    "Non guardare il questionario GINO: l'annotazione è in cieco. La verità è il protocollo.",
    "Non confrontarti con l'altro annotatore prima di aver finito.",
]


def scrivi(path, p, e):
    mancanti = pd.DataFrame(columns=["codice", "nome", "cat_principale", "cat_secondarie", "note"])
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame({"istruzioni": ISTR}).to_excel(w, sheet_name="istruzioni", index=False)
        p.to_excel(w, sheet_name="protocolli", index=False)
        cat[["codice", "etichetta"]].to_excel(w, sheet_name="categorie", index=False)
        e.to_excel(w, sheet_name="enti_AI", index=False)
        mancanti.to_excel(w, sheet_name="enti_mancanti", index=False)
    wb = load_workbook(path)
    for ws in wb:
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        ws.freeze_panes = "A2"
    wb["istruzioni"].column_dimensions["A"].width = 150
    for row in wb["istruzioni"].iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True)
    ws = wb["enti_AI"]
    # id A, codice B, regione C, annotatori D, nome_AI E, tipo_AI F, cat_AI G, giudizio H, nome_corretto I, cat_principale J, cat_secondarie K, note L
    for col, wd in zip("ABCDEFGHIJKL", [6, 8, 16, 10, 55, 40, 8, 14, 35, 14, 14, 40]):
        ws.column_dimensions[col].width = wd
    n = max(ws.max_row, 2)
    dv_g = DataValidation(type="list", formula1=f'"{GIUDIZI}"', allow_blank=True)
    dv_c = DataValidation(type="whole", operator="between", formula1="1", formula2="30", allow_blank=True)
    ws.add_data_validation(dv_g); ws.add_data_validation(dv_c)
    dv_g.add(f"H2:H{n}"); dv_c.add(f"J2:J{n}")
    wm = wb["enti_mancanti"]
    dv_m = DataValidation(type="whole", operator="between", formula1="1", formula2="30", allow_blank=True)
    wm.add_data_validation(dv_m); dv_m.add("C2:C300")
    for col, wd in zip("ABCDE", [8, 55, 14, 14, 40]):
        wm.column_dimensions[col].width = wd
    wb["categorie"].column_dimensions["B"].width = 100
    wb["protocolli"].column_dimensions["D"].width = 90
    wb["protocolli"].column_dimensions["E"].width = 50
    wb.save(path)


scrivi(OUT / "gold_enti_v1_template.xlsx", PROT.copy(), ENTI.copy())
for x in ("A1", "A2"):
    m = PROT.annotatori.isin([x, "A1+A2"])
    scrivi(OUT / f"gold_enti_v1_{x}.xlsx", PROT[m].copy(), ENTI[ENTI.codice.isin(PROT.codice[m])].copy())
    print(x, int(m.sum()), "protocolli", int(ENTI.codice.isin(PROT.codice[m]).sum()), "enti")
print("template:", len(PROT), "protocolli", len(ENTI), "enti")
