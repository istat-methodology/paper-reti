"""Revisione manuale degli enti del nucleo (stadio 1).

  prepara:  python revisione_enti.py prepara --tag v1_2_testo [--codici 0901,0902] [--n-run 5] [--campione 0.1]
            -> revisione/revisione_<tag>.xlsx (enti da rivedere, con colonne di decisione)
               revisione/enti_automatici_<tag>.csv (tutti gli enti, prima della revisione)
  applica:  python revisione_enti.py applica --tag v1_2_testo --decisioni revisione/revisione_<tag>_compilato.xlsx
            -> revisione/enti_finali_<tag>.csv (decisioni applicate)

Un ente è da rivedere se: non è presente in tutte le run, cambia ruolo tra run, ha un nome simile a un altro ente dello
stesso protocollo (possibile doppione), ha la citazione non trovata nel testo, oppure è nel campione casuale degli enti
unanimi (per stimare la precisione del consenso). Le decisioni sono righe (protocollo, nome) -> azione: restano
valide se si rilanciano le run, perché si agganciano a qualsiasi variante del nome."""
import argparse, csv, difflib, json, random, sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent)); sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from config import REPO  # noqa: E402
from stabilita_tollerante import BASE, formato_v2, norm, uguali  # noqa: E402

OUT = REPO / "revisione"
FORZA = {"firmatario": 3, "componente_rete": 2, "solo_citato": 1}
AZIONI = ["ok", "elimina", "unisci"]


def carica(tag, codici=None, n_run=None):
    """{codice: (n_run, [gruppo])}; gruppo = nomi (Counter), varianti (norm), run, ruoli (Counter), citazioni"""
    runs = {}
    for f in sorted((BASE / tag).glob("[0-9]*_run*.json")):
        d = formato_v2(json.loads(f.read_text(encoding="utf-8"))); c = d["meta"]["codice"]
        if codici and c not in codici:
            continue
        runs.setdefault(c, {})[d["meta"]["run"]] = d["enti_identificati"]
    out = {}
    for c, r in sorted(runs.items()):
        ids = sorted(r)[:n_run] if n_run else sorted(r)
        gr = []
        for i in ids:
            for e in r[i]:
                k = norm(e["nome"])
                if not k:
                    continue
                g = next((g for g in gr if any(uguali(k, v) for v in g["varianti"])), None)
                if g is None:
                    g = dict(nomi=Counter(), varianti=set(), run=set(), ruoli=Counter(), cit=[], cit_ok=Counter())
                    gr.append(g)
                g["varianti"].add(k)
                if i not in g["run"]:
                    g["ruoli"][e["ruolo_nel_documento"]] += 1
                g["run"].add(i); g["nomi"][e["nome"]] += 1
                g["cit"].append((FORZA[e["ruolo_nel_documento"]], e["citazione"]))
                g["cit_ok"][bool(e.get("citazione_trovata_nel_testo", True))] += 1
        out[c] = (len(ids), gr)
    return out


def nome(g):
    return max(g["nomi"], key=lambda n: (g["nomi"][n], len(n)))


def prepara(a):
    random.seed(20261031)
    dati = carica(a.tag, a.codici.split(",") if a.codici else None, a.n_run)
    OUT.mkdir(exist_ok=True)
    righe, tutti, coppia_id = [], [], 0
    for c, (n, gr) in dati.items():
        coppie = {}
        for i in range(len(gr)):
            for j in range(i + 1, len(gr)):
                x, y = nome(gr[i]), nome(gr[j]); nx, ny = norm(x), norm(y)
                sx, sy = set(nx.split()), set(ny.split())
                sub = min(len(sx), len(sy)) >= 2 and (sx <= sy or sy <= sx)   # un nome è contenuto nell'altro
                if sub or difflib.SequenceMatcher(None, nx, ny).ratio() >= 0.9:
                    coppia_id += 1; coppie.setdefault(i, []).append(coppia_id); coppie.setdefault(j, []).append(coppia_id)
        for idx, g in enumerate(gr):
            motivi = []
            if len(g["run"]) < n: motivi.append(f"non in tutte le run ({len(g['run'])}/{n})")
            if len(g["ruoli"]) > 1: motivi.append("ruolo instabile")
            if idx in coppie: motivi.append("possibile doppione (coppia " + ",".join(map(str, coppie[idx])) + ")")
            if g["cit_ok"][False] > g["cit_ok"][True]: motivi.append("citazione non trovata nel testo")
            if not motivi and random.random() < a.campione: motivi.append("campione dei presenti in tutte le run")
            ruolo = g["ruoli"].most_common(1)[0][0]
            cit = max(g["cit"])[1]
            riga = dict(codice=c, nome=nome(g), n_run_presente=len(g["run"]), n_run=n, ruolo_prevalente=ruolo,
                        ruoli="; ".join(f"{k} {v}" for k, v in g["ruoli"].most_common()),
                        varianti_nome=" | ".join(sorted(g["nomi"])), motivo=" + ".join(motivi), citazione=cit.replace("\n", " ")[:300])
            tutti.append({k: riga[k] for k in ("codice", "nome", "n_run_presente", "n_run", "ruolo_prevalente", "ruoli", "varianti_nome")})
            if motivi:
                righe.append(riga)
    righe.sort(key=lambda r: (r["codice"], -FORZA[r["ruolo_prevalente"]], r["nome"]))
    with open(OUT / f"enti_automatici_{a.tag}.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(tutti[0]), delimiter=";"); w.writeheader(); w.writerows(tutti)
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ws = wb.active; ws.title = "da_rivedere"
    cols = list(righe[0]) + ["azione", "unisci_con", "nome_finale", "ruolo_finale", "note"]
    ws.append(cols)
    for i, r in enumerate(righe, 2):
        ws.append([r[k] for k in list(r)] + [None] * 5)
    dv = DataValidation(type="list", formula1='"' + ",".join(AZIONI) + '"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"{chr(65 + cols.index('azione'))}2:{chr(65 + cols.index('azione'))}{len(righe) + 1}")
    dv2 = DataValidation(type="list", formula1='"firmatario,componente_rete,solo_citato"', allow_blank=True); ws.add_data_validation(dv2)
    dv2.add(f"{chr(65 + cols.index('ruolo_finale'))}2:{chr(65 + cols.index('ruolo_finale'))}{len(righe) + 1}")
    ws.freeze_panes = "C2"
    for col, w_ in zip("ABCDEFGHIJKLMNO", [8, 45, 9, 6, 16, 28, 45, 40, 60, 10, 30, 30, 16, 30]): ws.column_dimensions[col].width = w_
    ist = wb.create_sheet("istruzioni", 0)
    for t in ["Revisione degli enti del nucleo. Una riga per ente; se non scrivi nulla in 'azione' l'ente resta com'è.",
              "azione = ok (confermato), elimina (non è un ente della rete o non è un ente), unisci (stesso ente di un'altra riga dello stesso protocollo: scrivi il suo nome in 'unisci_con').",
              "nome_finale, ruolo_finale: solo se vuoi cambiare il nome o il ruolo.",
              "Il motivo dice perché la riga è qui. 'campione dei presenti in tutte le run' serve a stimare quanti errori stabili ci sono: giudicale come le altre.",
              "Ruolo: firmatario = sottoscrive; componente_rete = parte della rete o con impegni; solo_citato = nominato senza essere parte della rete."]:
        ist.append([t])
    ist.column_dimensions["A"].width = 160
    wb.save(OUT / f"revisione_{a.tag}.xlsx")
    c = Counter(m.split(" (")[0].split(" +")[0] for r in righe for m in r["motivo"].split(" + "))
    print(f"{len(dati)} protocolli, {len(tutti)} enti, {len(righe)} da rivedere ({len(righe) / len(dati):.1f} per protocollo)")
    for k, v in c.most_common(): print(f"  {k}: {v}")
    print("->", OUT / f"revisione_{a.tag}.xlsx")


def applica(a):
    import openpyxl
    dati = carica(a.tag, None, a.n_run)
    ws = openpyxl.load_workbook(a.decisioni)["da_rivedere"]
    h = [c.value for c in ws[1]]; dec = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(h, r))
        if d.get("azione"):
            for v in str(d["varianti_nome"] or d["nome"]).split(" | "):
                dec[(str(d["codice"]), norm(v))] = d
    finali = []
    for c, (n, gr) in dati.items():
        eliminati, uniti = set(), {}
        info = {}
        for idx, g in enumerate(gr):
            d = next((dec[(c, v)] for v in g["varianti"] if (c, v) in dec), None); info[idx] = d
            if d and d["azione"] == "elimina": eliminati.add(idx)
            if d and d["azione"] == "unisci" and d.get("unisci_con"):
                t = norm(d["unisci_con"])
                tgt = next((j for j, h2 in enumerate(gr) if j != idx and any(uguali(t, v) for v in h2["varianti"])), None)
                if tgt is not None: uniti[idx] = tgt
        for idx, g in enumerate(gr):
            if idx in eliminati or idx in uniti: continue
            run = set(g["run"]); ruoli = Counter(g["ruoli"])
            for src, tgt in uniti.items():
                if tgt == idx: run |= gr[src]["run"]; ruoli += gr[src]["ruoli"]
            d = info[idx]
            nm = (d.get("nome_finale") if d else None) or nome(g)
            ruolo = (d.get("ruolo_finale") if d else None) or ruoli.most_common(1)[0][0]
            finali.append(dict(codice=c, nome=nm, ruolo=ruolo, n_run_presente=len(run), n_run=n,
                               stato="rivisto" if d or any(t == idx for t in uniti.values()) else "automatico"))
    with open(OUT / f"enti_finali_{a.tag}.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(finali[0]), delimiter=";"); w.writeheader(); w.writerows(finali)
    print(f"{len(finali)} enti finali ({sum(x['stato'] == 'rivisto' for x in finali)} rivisti) -> {OUT / ('enti_finali_' + a.tag + '.csv')}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("prepara"); p.add_argument("--tag", required=True); p.add_argument("--codici"); p.add_argument("--n-run", type=int); p.add_argument("--campione", type=float, default=0.1)
    q = sp.add_parser("applica"); q.add_argument("--tag", required=True); q.add_argument("--decisioni", required=True); q.add_argument("--n-run", type=int)
    a = ap.parse_args(); prepara(a) if a.cmd == "prepara" else applica(a)
